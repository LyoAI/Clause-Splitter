"""Extract text from documents (PDF / DOCX / DOC / XLSX / TXT / images)."""

from __future__ import annotations

from pathlib import Path


def extract_text(file_path: str | Path, max_chars: int = 20000) -> str:
    """Auto-detect file type and extract plain text.

    Args:
        file_path: Path to the document.
        max_chars: Maximum characters to extract.

    Returns:
        Extracted plain text, or an error placeholder string.
    """
    path = Path(file_path)
    if not path.exists():
        return ""

    suffix = path.suffix.lower()

    try:
        if suffix == ".pdf":
            return _parse_pdf(path, max_chars)
        if suffix == ".docx":
            return _parse_docx(path, max_chars)
        if suffix == ".doc":
            return _parse_doc(path, max_chars)
        if suffix in (".xlsx", ".xls"):
            return _parse_xlsx(path, max_chars)
        if suffix in (".png", ".jpg", ".jpeg", ".bmp", ".tiff"):
            return f"[Image] {path.name} — OCR not available, manual review required."
        # Plain text fallback
        return path.read_text(encoding="utf-8", errors="replace")[:max_chars]
    except Exception as e:
        return f"[Parse error: {e}]"


def _parse_pdf(path: Path, max_chars: int) -> str:
    """Extract text from PDF files using pdfplumber."""
    import pdfplumber
    texts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                texts.append(t)
            if sum(len(x) for x in texts) >= max_chars:
                break
    return "\n".join(texts)[:max_chars]


def _parse_docx(path: Path, max_chars: int) -> str:
    """Extract text from .docx (Office Open XML) files."""
    from docx import Document
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)[:max_chars]


def _parse_doc(path: Path, max_chars: int) -> str:
    """Extract text from legacy .doc (binary) files.

    Tries multiple backends since old-format .doc is not supported by python-docx.
    """
    import subprocess

    # Strategy 1: antiword (common on Linux/macOS)
    try:
        result = subprocess.run(
            ["antiword", str(path)],
            capture_output=True, text=True, timeout=30,
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout.strip()[:max_chars]
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass

    # Strategy 2: textract (if installed)
    try:
        import textract
        text = textract.process(str(path)).decode("utf-8", errors="replace")
        if text.strip():
            return text.strip()[:max_chars]
    except Exception:
        pass

    # Strategy 3: python-docx — works for misnamed .docx files
    try:
        from docx import Document
        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        if paragraphs:
            return "\n".join(paragraphs)[:max_chars]
    except Exception:
        pass

    return (
        f"[Legacy .doc] {path.name} — unable to extract text. "
        f"Please convert to .docx or .pdf."
    )


def _parse_xlsx(path: Path, max_chars: int) -> str:
    """Extract text from Excel files (all sheets, tab-separated)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        lines.append(f"[Sheet: {sheet}]")
        for row in ws.iter_rows(values_only=True):
            row_texts = [str(v) for v in row if v is not None]
            if row_texts:
                lines.append("\t".join(row_texts))
    return "\n".join(lines)[:max_chars]
