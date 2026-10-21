"""Tests for the Excel multi-regulation container (``containers.split_workbook``).

An internal-regulation library is often one Excel where every row is an article
(第一条 / 第二条 …) and rows are grouped by a 制度名称 column. ``split_workbook``
splits such a workbook into per-regulation ``RegulationDoc`` block lists so the
rest of the ingest pipeline treats each 制度 as its own document, chunked per-条
with the 条号 as ``meta.section``.

Covers:
- Per-制度 split with deterministic (no-LLM) column mapping via header keywords.
- Chunking: one chunk per 条, ``meta.section`` = 条号.
- Blank 制度名称 rows → the ``_未命名文档_`` fallback doc.
- Unrecognizable workbook (no content column) → clear ``ValueError``.
- End-to-end via ``ingest``: the store is keyed by 制度名.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import openpyxl
import pytest

from chillrag.tools.parser import chunk_doc

from clause_splitter.compare.containers import split_workbook
from clause_splitter.compare.headings import infer_regulation_headings
from clause_splitter.compare.ingest import ingest


def _break_dimension_ref(path: Path) -> None:
    """把首个 worksheet 的 <dimension> 改成 A1,复现系统导出文件的 read_only 漏行坑。"""
    zpath = Path(path)
    tmp = zpath.with_suffix(".dim.tmp.xlsx")
    with zipfile.ZipFile(zpath) as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}
    sheet = next(n for n in names if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n))
    xml = data[sheet].decode("utf-8")
    xml = re.sub(r"<dimension[^>]*/>", '<dimension ref="A1"/>', xml, count=1)
    data[sheet] = xml.encode("utf-8")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    tmp.replace(zpath)


def _write_workbook(path: Path, rows: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "制度"
    ws.append(["制度名称", "条款名称", "条款内容"])
    for r in rows:
        ws.append([r.get("doc", ""), r.get("id", ""), r.get("text", "")])
    wb.save(path)


# ------------------------------------------------------------- splitting -----

def test_split_workbook_per_regulation_articles(tmp_path: Path) -> None:
    wb_path = tmp_path / "内规库.xlsx"
    _write_workbook(
        wb_path,
        [
            {"doc": "网络数据安全管理条例", "id": "第一条", "text": "为了规范网络数据处理活动。", },
            {"doc": "网络数据安全管理条例", "id": "第二条", "text": "适用对象为本公司。", },
            {"doc": "个人信息保护办法", "id": "第一条", "text": "个人信息处理应当遵循合法原则。", },
            {"doc": "个人信息保护办法", "id": "第二条", "text": "不得非法收集他人信息。", },
        ],
    )

    docs = split_workbook(wb_path)
    assert [d.doc_id for d in docs] == ["网络数据安全管理条例", "个人信息保护办法"]
    assert all(d.source_file == str(wb_path) for d in docs)

    # Each regulation → alternating 条号 title (level 1) + body paragraph blocks.
    first = docs[0]
    assert [(b.type, b.text) for b in first.blocks] == [
        ("title", "第一条"),
        ("paragraph", "为了规范网络数据处理活动。"),
        ("title", "第二条"),
        ("paragraph", "适用对象为本公司。"),
    ]


def test_container_blocks_chunk_per_article(tmp_path: Path) -> None:
    wb_path = tmp_path / "内规库.xlsx"
    _write_workbook(
        wb_path,
        [
            {"doc": "数据安全制度", "id": "第一条", "text": "第一条　数据分类分级。"},
            {"doc": "数据安全制度", "id": "第二条", "text": "数据由专人管理。"},
        ],
    )
    reg = split_workbook(wb_path)[0]
    chunks = chunk_doc(reg.doc_id, infer_regulation_headings(reg.blocks), level=4, min_tokens=0)
    assert [c.meta["section"] for c in chunks] == ["第一条", "第二条"]
    assert len(chunks) == 2
    # The cell already leads with "第一条　" — the label is not duplicated.
    assert chunks[0].text == "第一条\n\n数据分类分级。"


def test_blank_doc_name_falls_back_to_unnamed(tmp_path: Path) -> None:
    wb_path = tmp_path / "内规库.xlsx"
    _write_workbook(
        wb_path,
        [{"id": "第一条", "text": "没有制度名的条款。"}],
    )
    docs = split_workbook(wb_path)
    assert len(docs) == 1
    assert docs[0].doc_id == "_未命名文档_"
    assert docs[0].blocks[0].text == "第一条"


def test_unrecognizable_workbook_raises(tmp_path: Path) -> None:
    wb_path = tmp_path / "随机.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "部门"])
    ws.append(["张三", "IT"])
    wb.save(wb_path)
    with pytest.raises(ValueError, match="无法识别"):
        split_workbook(wb_path)


# ------------------------------------------------- internal-network layouts ---
# 内网模板的常见错位:表头不在第 1 行、数据不在第一个 sheet、制度名列的表头
# 用自定义名字(关键词覆盖不到)→ 这些用例保证 split_workbook 仍能按制度分组。


def test_header_row_after_title_rows(tmp_path: Path) -> None:
    """表头在第 3 行(前面是合并标题 + 编制信息行)也能被定位。"""
    wb_path = tmp_path / "内规库.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["XX公司内部制度汇编"])            # row 1: 合并标题
    ws.append(["编制部门:合规部  日期:2026-01-01"])  # row 2: 说明行
    ws.append(["制度名称", "条款名称", "条款内容"])  # row 3: 真实表头
    ws.append(["网络数据安全管理条例", "第一条", "为了规范网络数据处理活动。"])
    ws.append(["网络数据安全管理条例", "第二条", "适用对象为本公司。"])
    ws.append(["个人信息保护办法", "第一条", "个人信息处理应当遵循合法原则。"])
    wb.save(wb_path)

    docs = split_workbook(wb_path)
    assert [d.doc_id for d in docs] == ["网络数据安全管理条例", "个人信息保护办法"]
    assert [(b.type, b.text) for b in docs[0].blocks] == [
        ("title", "第一条"), ("paragraph", "为了规范网络数据处理活动。"),
        ("title", "第二条"), ("paragraph", "适用对象为本公司。"),
    ]


def test_data_on_second_sheet_skips_cover(tmp_path: Path) -> None:
    """第一个 sheet 是封面(无表头关键词),第二个 sheet 才是数据表 → 自动选数据 sheet。"""
    wb_path = tmp_path / "内规库.xlsx"
    wb = openpyxl.Workbook()
    cover = wb.active
    cover.title = "封面"
    cover.append(["XX公司制度汇编说明"])
    cover.append(["仅供内部使用"])
    data = wb.create_sheet("制度明细")
    data.append(["内部制度名称", "条款内容"])
    data.append(["数据安全管理办法", "第一条　数据分类分级。"])
    data.append(["数据安全管理办法", "第二条　数据由专人管理。"])
    wb.save(wb_path)

    docs = split_workbook(wb_path)
    assert [d.doc_id for d in docs] == ["数据安全管理办法"]
    assert len(docs[0].blocks) == 4  # 条号 title + body × 2 条


def test_custom_doc_name_header_detected_by_content(tmp_path: Path) -> None:
    """制度名列表头叫「归属制度」(关键词未覆盖)→ 按"短值重复"从内容推断分组。"""
    wb_path = tmp_path / "内规库.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["归属制度", "条款内容"])
    ws.append(["制度甲", "第一条　甲内容一。"])
    ws.append(["制度甲", "第二条　甲内容二。"])
    ws.append(["制度乙", "第一条　乙内容一。"])
    wb.save(wb_path)

    docs = split_workbook(wb_path)
    assert [d.doc_id for d in docs] == ["制度甲", "制度乙"]
    assert len(docs[0].blocks) == 4
    assert len(docs[1].blocks) == 2


def test_real_layout_ignores_full_doc_text_column(tmp_path: Path) -> None:
    """内网真实模板:每行含「制度内容」(整篇全文,每行重复)+「条款内容」。

    期望:「制度内容」不进任何 chunk(否则每块都≈整篇制度),doc_id 不带双重《》,
    检索标签用「条款名称」(第一节/第二节),每行 = 一条。
    """
    wb_path = tmp_path / "内规库.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["内部制度名称", "制度内容", "IOBS_KEY", "章节", "条款名称", "条款内容"])
    full_a = "【完整制度全文】第一条　外包人员的准入条件…第二条　须签署保密协议…（整篇重复）"
    full_b = "【完整制度全文】第一条　数据分级分类…（整篇重复）"
    ws.append(["《IT人力外包管理办法》", full_a, "K001", "第一章", "第一节", "第一条　外包人员的准入条件…"])
    ws.append(["《IT人力外包管理办法》", full_a, "K002", "第一章", "第二节", "第二条　外包人员须签署保密协议…"])
    ws.append(["《数据安全管理细则》", full_b, "K003", "第一章", "第一节", "第一条　数据分级分类…"])
    wb.save(wb_path)

    docs = split_workbook(wb_path)
    # doc_id 不再出现《《…》》双书名号。
    assert [d.doc_id for d in docs] == ["IT人力外包管理办法", "数据安全管理细则"]

    reg = docs[0]
    assert [(b.type, b.text) for b in reg.blocks] == [
        ("title", "第一节"), ("paragraph", "第一条　外包人员的准入条件…"),
        ("title", "第二节"), ("paragraph", "第二条　外包人员须签署保密协议…"),
    ]
    chunks = chunk_doc(reg.doc_id, infer_regulation_headings(reg.blocks), level=4, min_tokens=0)
    # 条款名称(节)作为 section 父级,正文里的真实条号(第一条)作为叶子。
    assert [c.meta["section"] for c in chunks] == ["第一节 > 第一条", "第二节 > 第二条"]
    # 制度内容(整篇全文)绝不能进 chunk。
    assert all("【完整制度全文】" not in c.text for c in chunks)
    assert chunks[0].text == "第一条\n\n外包人员的准入条件…"
    assert chunks[1].text == "第二条\n\n外包人员须签署保密协议…"


def test_broken_dimension_still_reads_all_columns(tmp_path: Path) -> None:
    """系统导出的 xlsx 常带不完整的 <dimension>;read_only=True 流式只吐 A1 一格。

    回归:预览与抽取都必须用 read_only=False 物化全部行,否则表头/数据读不全,
    直接报"可用表头:内部制度名称"。
    """
    wb_path = tmp_path / "内规库.xlsx"
    _write_workbook(
        wb_path,
        [
            {"doc": "网络数据安全管理条例", "id": "第一条", "text": "为了规范网络数据处理活动。"},
            {"doc": "网络数据安全管理条例", "id": "第二条", "text": "适用对象为本公司。"},
        ],
    )
    _break_dimension_ref(wb_path)

    docs = split_workbook(wb_path)
    assert [d.doc_id for d in docs] == ["网络数据安全管理条例"]
    assert [(b.type, b.text) for b in docs[0].blocks] == [
        ("title", "第一条"), ("paragraph", "为了规范网络数据处理活动。"),
        ("title", "第二条"), ("paragraph", "适用对象为本公司。"),
    ]


def test_blank_article_id_inferred_from_text(tmp_path: Path) -> None:
    wb_path = tmp_path / "内规库.xlsx"
    _write_workbook(
        wb_path,
        [{"doc": "某制度", "text": "第三条　明确责任分工。"}],  # 条款名称 blank
    )
    reg = split_workbook(wb_path)[0]
    assert reg.blocks[0].text == "第三条"
    assert reg.blocks[1].text == "明确责任分工。"


# ------------------------------------------------------------- end-to-end ----

def test_ingest_routes_workbook_to_per_doc_store(tmp_path: Path) -> None:
    wb_path = tmp_path / "内规库.xlsx"
    _write_workbook(
        wb_path,
        [
            {"doc": "制度甲", "id": "第一条", "text": "甲第一条内容。"},
            {"doc": "制度乙", "id": "第一条", "text": "乙第一条内容。"},
        ],
    )
    store = ingest([wb_path], name="tmp", persist=False)
    assert set(store.docs) == {"制度甲", "制度乙"}
    chunks = store.all_chunks()
    assert [c.meta["section"] for c in chunks] == ["第一条", "第一条"]
    # Provenance: per-doc source entry, kind=excel_regulation, sha256 of the workbook.
    assert set(store.sources) == {"制度甲", "制度乙"}
    for entry in store.sources.values():
        assert entry["kind"] == "excel_regulation"
        assert entry["source_file"] == str(wb_path)
        assert len(entry["sha256"]) == 64
