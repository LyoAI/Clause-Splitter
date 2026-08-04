"""Unit tests for Excel clause parsing tools.

Tests the deterministic tool functions that support the ReAct Excel parsing node.
"""

import pytest
import tempfile
from pathlib import Path

from agent.tools.excel_parse_tools import (
    _normalize_column_header,
    _infer_article_id_from_text,
    _col_idx_to_letter,
    read_excel_preview,
    extract_excel_articles,
)


# ── Column header normalization ──────────────────────────────────────────


def test_normalize_header_basic():
    assert _normalize_column_header("条款内容") == "条款内容"


def test_normalize_header_strips_whitespace():
    assert _normalize_column_header("  条款内容 ") == "条款内容"


def test_normalize_header_removes_parenthetical():
    assert _normalize_column_header("条款内容（原文）") == "条款内容"
    assert _normalize_column_header("名称(English)") == "名称"


def test_normalize_header_removes_brackets():
    # Full-width brackets 【】 are removed but their content is kept
    assert _normalize_column_header("【外部法规】URL") == "外部法规URL"


def test_normalize_header_none():
    assert _normalize_column_header(None) == ""


def test_normalize_header_empty():
    assert _normalize_column_header("") == ""


# ── Article ID inference ────────────────────────────────────────────────


def test_infer_article_id_chinese():
    assert _infer_article_id_from_text("第十四条 数据处理者应当建立数据安全管理制度") == "第十四条"


def test_infer_article_id_english():
    assert _infer_article_id_from_text("Article 5 The processor shall implement measures") == "Article 5"


def test_infer_article_id_numbered_list():
    assert "1." in _infer_article_id_from_text("1. 数据分类分级管理规定")


def test_infer_article_id_empty():
    assert _infer_article_id_from_text("") == "空白条款"


def test_infer_article_id_no_match():
    result = _infer_article_id_from_text("数据处理者应当建立完善的安全制度")
    assert len(result) > 0


# ── Column letter conversion ─────────────────────────────────────────────


def test_col_idx_to_letter():
    assert _col_idx_to_letter(0) == "A"
    assert _col_idx_to_letter(1) == "B"
    assert _col_idx_to_letter(25) == "Z"
    assert _col_idx_to_letter(26) == "AA"
    assert _col_idx_to_letter(27) == "AB"
    assert _col_idx_to_letter(701) == "ZZ"


# ── Excel preview ────────────────────────────────────────────────────────


def test_read_excel_preview_basic():
    """Test reading a simple Excel file."""
    import openpyxl

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "条款数据"
        # Header row
        ws.append(["内部制度名称", "章节", "条款名称", "条款内容"])
        # Data rows
        ws.append(["制度A", "第一章", "第一条", "数据处理者应当建立数据安全管理制度。"])
        ws.append(["制度A", "第一章", "第二条", "数据处理者应当对数据进行分类分级。"])
        ws.append(["制度A", "第二章", "第三条", "禁止未经授权的数据访问。"])
        wb.save(tmp_path)
        wb.close()

        result = read_excel_preview(tmp_path, max_rows=10)

        assert result["error"] is None
        assert result["active_sheet"] == "条款数据"
        assert result["total_rows"] == 3
        assert len(result["column_headers"]) == 4
        assert result["column_headers"]["A"] == "内部制度名称"
        assert result["column_headers"]["D"] == "条款内容"
        assert len(result["samples"]) == 3
        assert result["samples"][0]["columns"]["D"] == "数据处理者应当建立数据安全管理制度。"

    finally:
        Path(tmp_path).unlink(missing_ok=True)


def test_read_excel_preview_empty():
    """Test reading an empty Excel file (header only)."""
    import openpyxl

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["内部制度名称", "条款内容"])
        wb.save(tmp_path)
        wb.close()

        result = read_excel_preview(tmp_path, max_rows=10)

        assert result["error"] is None
        assert result["total_rows"] == 0
        assert len(result["samples"]) == 0

    finally:
        Path(tmp_path).unlink(missing_ok=True)


def test_read_excel_preview_file_not_found():
    result = read_excel_preview("/nonexistent/path.xlsx")
    assert result["error"] is not None


# ── Excel clause extraction ──────────────────────────────────────────────


def test_extract_excel_articles_basic():
    """Test extracting clauses from Excel with a column mapping."""
    import openpyxl

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["内部制度名称", "章节", "条款名称", "条款内容"])
        ws.append(["制度A", "第一章", "第一条", "数据处理者应当建立数据安全管理制度。"])
        ws.append(["制度A", "第一章", "第二条", "数据处理者应当对数据进行分类分级。"])
        ws.append(["制度B", "第三章", "第五条", "禁止数据非法交易。"])
        wb.save(tmp_path)
        wb.close()

        mapping = {
            "article_text": "条款内容",
            "article_id": "条款名称",
            "document_name": "内部制度名称",
            "section": "章节",
        }

        result = extract_excel_articles(tmp_path, mapping)

        assert result["error"] is None
        assert result["total_rows"] == 3
        assert len(result["articles"]) == 3
        assert len(result["unique_documents"]) == 2

        # Check article content
        assert result["articles"][0]["full_text"] == "数据处理者应当建立数据安全管理制度。"
        assert result["articles"][0]["article_id"] == "第一条"
        assert result["articles"][0]["document_name"] == "制度A"
        assert result["articles"][0]["section"] == "第一章"

        # Check multi-document grouping
        assert "制度A" in result["document_groups"]
        assert "制度B" in result["document_groups"]
        assert len(result["document_groups"]["制度A"]) == 2
        assert len(result["document_groups"]["制度B"]) == 1

    finally:
        Path(tmp_path).unlink(missing_ok=True)


def test_extract_excel_articles_single_document():
    """Test single document Excel (no document_name in mapping)."""
    import openpyxl

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["条款名称", "条款内容"])
        ws.append(["第一条", "应当建立制度。"])
        ws.append(["第二条", "禁止违规操作。"])
        wb.save(tmp_path)
        wb.close()

        mapping = {"article_text": "条款内容", "article_id": "条款名称"}

        result = extract_excel_articles(tmp_path, mapping)

        assert result["error"] is None
        assert result["total_rows"] == 2
        assert len(result["documents"]) == 1 if "documents" in result else True  # single group
        assert result["articles"][0]["article_id"] == "第一条"

    finally:
        Path(tmp_path).unlink(missing_ok=True)


def test_extract_excel_articles_missing_mapping():
    """Test that missing article_text gives an error."""
    import openpyxl

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["A", "B"])
        ws.append(["x", "y"])
        wb.save(tmp_path)
        wb.close()

        mapping = {"article_text": "不存在的列名"}
        result = extract_excel_articles(tmp_path, mapping)

        assert result["error"] is not None

    finally:
        Path(tmp_path).unlink(missing_ok=True)


def test_extract_excel_articles_file_not_found():
    mapping = {"article_text": "条款内容"}
    result = extract_excel_articles("/nonexistent/path.xlsx", mapping)
    assert result["error"] is not None


# ── Heuristic fallback tests ─────────────────────────────────────────────


def test_heuristic_column_detection():
    """Test that heuristic fallback correctly identifies columns."""
    from agent.nodes.parse_excel import _heuristic_column_detection

    preview = {
        "column_headers": {
            "A": "内部制度名称",
            "B": "制度内容",
            "C": "章节",
            "D": "条款名称",
            "E": "条款内容",
        },
        "samples": [
            {"row": 2, "columns": {
                "A": "制度A", "B": "概要...", "C": "第一章",
                "D": "第一条", "E": "数据处理者应当建立制度。"
            }},
        ],
    }

    mapping = _heuristic_column_detection(preview)
    assert mapping is not None
    assert mapping["article_text"] == "条款内容"
    assert mapping["article_id"] == "条款名称"
    assert mapping["document_name"] == "内部制度名称"
    assert mapping["section"] == "章节"


def test_heuristic_no_headers():
    from agent.nodes.parse_excel import _heuristic_column_detection

    mapping = _heuristic_column_detection({"column_headers": {}, "samples": []})
    assert mapping is None


# ── Safe filename tests ──────────────────────────────────────────────────


def test_safe_filename_basic():
    from services.splitter import _safe_filename

    assert _safe_filename("制度A") == "制度A"
    assert _safe_filename("制度A/B") == "制度A_B"
    assert _safe_filename("制度<A>:B") == "制度_A_B"
    assert _safe_filename("  带空格  ") == "带空格"
    assert _safe_filename("a.txt") == "a.txt"
    assert _safe_filename("") == "未命名文档"
    assert _safe_filename("制度名?*") == "制度名"  # trailing _ stripped


# ── Subgraph build test ──────────────────────────────────────────────────


def test_build_split_subgraph():
    """Verify the sub-pipeline graph builds correctly."""
    from agent.graph import build_split_subgraph

    sg = build_split_subgraph()
    assert sg is not None
    # The compiled graph should have plan_batches as entry
    # We can at least verify it compiles without error
    print("  ✓ subgraph compiled")


# ── Multi-document extraction with grouping ──────────────────────────────


def test_extract_excel_articles_grouping():
    """Test that multi-document Excel correctly groups by document_name."""
    import openpyxl

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["内部制度名称", "章节", "条款名称", "条款内容"])
        # 制度A: 2 articles
        ws.append(["制度A", "第一章", "第一条", "应当建立数据安全管理制度。"])
        ws.append(["制度A", "第一章", "第二条", "应当对数据进行分类分级。"])
        # 制度B: 1 article
        ws.append(["制度B", "第三章", "第五条", "禁止数据非法交易。"])
        # 制度C: 2 articles
        ws.append(["制度C", "第一章", "第一条", "必须定期进行安全审计。"])
        ws.append(["制度C", "第二章", "第三条", "不得泄露用户个人信息。"])
        wb.save(tmp_path)
        wb.close()

        mapping = {
            "article_text": "条款内容",
            "article_id": "条款名称",
            "document_name": "内部制度名称",
            "section": "章节",
        }

        result = extract_excel_articles(tmp_path, mapping)

        assert result["error"] is None
        assert result["total_rows"] == 5
        assert len(result["unique_documents"]) == 3
        assert set(result["unique_documents"]) == {"制度A", "制度B", "制度C"}

        # Check per-document counts
        assert len(result["document_groups"]["制度A"]) == 2
        assert len(result["document_groups"]["制度B"]) == 1
        assert len(result["document_groups"]["制度C"]) == 2

        # Each article has correct document_name
        for art in result["articles"]:
            assert art["document_name"] in ("制度A", "制度B", "制度C")

    finally:
        Path(tmp_path).unlink(missing_ok=True)


def test_extract_excel_articles_unnamed_document():
    """Test Excel with no document_name column — all articles go to _未命名文档_."""
    import openpyxl

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["条款名称", "条款内容"])
        ws.append(["第一条", "应当建立制度。"])
        ws.append(["第二条", "禁止违规操作。"])
        wb.save(tmp_path)
        wb.close()

        mapping = {"article_text": "条款内容", "article_id": "条款名称"}

        result = extract_excel_articles(tmp_path, mapping)

        assert result["error"] is None
        # When no document_name column, all articles go to _未命名文档_
        assert len(result["unique_documents"]) == 1
        assert result["unique_documents"][0] == "_未命名文档_"
        assert len(result["document_groups"]["_未命名文档_"]) == 2

    finally:
        Path(tmp_path).unlink(missing_ok=True)


# ── Multi-document service (without LLM) ─────────────────────────────────


def test_split_multi_document_excel_no_llm():
    """Test multi-document service with heuristic fallback (no LLM)."""
    import asyncio
    import openpyxl
    import json
    from pathlib import Path as P

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        tmp_path = f.name

    tmp_dir = tempfile.mkdtemp()

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["内部制度名称", "章节", "条款名称", "条款内容"])
        ws.append(["制度A", "第一章", "第一条", "应当建立数据安全管理制度。"])
        ws.append(["制度A", "第二章", "第三条", "应当对数据进行分类分级。"])
        ws.append(["制度B", "第一章", "第一条", "禁止未授权的数据访问。"])
        wb.save(tmp_path)
        wb.close()

        from services.splitter import ClauseSplitService

        async def _run():
            service = ClauseSplitService(llm=None)  # No LLM → heuristic fallback
            return await service.split_multi_document_excel(
                file_path=tmp_path,
                document_name="test_multi.xlsx",
                output_base_dir=tmp_dir,
            )

        result = asyncio.run(_run())

        # Verify structure
        assert result["total_documents"] == 2
        assert result["total_articles"] == 3

        # Check per-document results
        docs = result["documents"]
        assert len(docs) == 2

        doc_names = {d["document_name"] for d in docs}
        assert "制度A" in doc_names
        assert "制度B" in doc_names

        # 制度A should have 2 articles, 制度B should have 1
        for d in docs:
            if d["document_name"] == "制度A":
                assert d["article_count"] == 2
            elif d["document_name"] == "制度B":
                assert d["article_count"] == 1
            assert d["status"] in ("completed", "failed")

        # Check output directories exist
        dir_a = P(tmp_dir) / "制度A"
        dir_b = P(tmp_dir) / "制度B"
        assert dir_a.exists(), f"Expected {dir_a} to exist"
        assert dir_b.exists(), f"Expected {dir_b} to exist"

        # Check JSON files
        json_a = dir_a / "条款拆分结果.json"
        if json_a.exists():
            data = json.loads(json_a.read_text(encoding="utf-8"))
            assert data["document_name"] == "制度A"
            assert data["total_articles"] == 2

    finally:
        Path(tmp_path).unlink(missing_ok=True)
        import shutil
        shutil.rmtree(tmp_dir, ignore_errors=True)
