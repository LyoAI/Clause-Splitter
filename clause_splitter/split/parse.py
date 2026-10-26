"""Excel 条款解析 —— 预览 / 抽取 / 表头归一 / 条款号推断.

列结构识别走关键词启发式 + 内容兜底,不调用 LLM。
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from clause_splitter.utils import get_logger

logger = get_logger()


# ── 启发式列检测 ───────────────────────────────────────────────────────

# Header keyword families: matched by substring after normalizing away
# parentheticals and whitespace.
_ARTICLE_TEXT_KEYWORDS = [
    "条款内容", "法规条款", "制度条款", "条文内容", "具体条款", "条款正文",
    "clause", "article_text", "content",
]
_ARTICLE_ID_KEYWORDS = [
    "条款名称", "条款编号", "条款标题", "条目名称", "条文号",
    "clause_name", "article_id", "title",
]
_DOC_NAME_KEYWORDS = [
    "内部制度名称", "制度名称", "外规名称", "法规名称", "文档名称", "文件名称",
    "document_name", "regulation_name",
]
_SECTION_KEYWORDS = ["章节", "所属章节", "章", "节", "section", "chapter"]
# Full-document columns repeat the whole regulation on every row — skip them
# entirely (as content or as context), or retrieval loses all granularity.
_FULL_TEXT_KEYWORDS = [
    "制度内容", "制度全文", "全文内容", "完整内容", "整篇内容",
    "full_document",
]
# Per-row context columns, prepended to the clause text.
_CONTEXT_KEYWORDS = [
    "适用条件", "适用范围", "适用对象", "备注",
    "document_context", "context",
]
_HEADER_KEYWORDS = (
    *_ARTICLE_TEXT_KEYWORDS,
    *_ARTICLE_ID_KEYWORDS,
    *_DOC_NAME_KEYWORDS,
    *_SECTION_KEYWORDS,
    *_FULL_TEXT_KEYWORDS,
    *_CONTEXT_KEYWORDS,
)


def _header_row_hits(cells: list[Any]) -> tuple[int, int]:
    """(表头关键词命中数, 非空去重单元格数) —— 用于定位真实表头行。"""
    kw_hits = 0
    seen: set[str] = set()
    for val in cells:
        if val is None:
            continue
        norm = _normalize_column_header(str(val))
        if not norm:
            continue
        seen.add(norm)
        low = norm.lower()
        if any(kw.lower() in low for kw in _HEADER_KEYWORDS):
            kw_hits += 1
    return kw_hits, len(seen)


def _find_header_row(ws, max_scan: int = 5) -> int:
    """真实表头行(1-based);找不到表头样行时回退到第 1 行。"""
    best_row = 1
    best_score = (-1, -1)
    for idx in range(1, max_scan + 1):
        values = next(ws.iter_rows(min_row=idx, max_row=idx, values_only=True), None)
        if values is None:
            break
        kw_hits, non_empty = _header_row_hits(values)
        score = (kw_hits, non_empty)
        if score > best_score:
            best_score = score
            best_row = idx
    if best_score == (0, 1) or best_score == (0, 0):
        return 1
    return best_row


def _match(col_name: str, keywords: list[str]) -> bool:
    """表头名(任意大小写/空格/括注)是否包含任一关键词。"""
    norm = _normalize_column_header(col_name).lower()
    return any(kw.lower() in norm for kw in keywords)


def _heuristic_column_detection(preview: dict) -> dict[str, str] | None:
    """无 LLM 时按关键词匹配表头(带内容兜底 + 制度名内容推断)。"""
    headers = preview.get("column_headers", {})
    if not headers:
        return None

    mapping: dict[str, str] = {}
    for col_letter, col_name in sorted(headers.items()):
        if not col_name:
            continue
        if _match(col_name, _FULL_TEXT_KEYWORDS):
            continue
        if not mapping.get("article_text") and _match(col_name, _ARTICLE_TEXT_KEYWORDS):
            mapping["article_text"] = col_name
        elif not mapping.get("article_id") and _match(col_name, _ARTICLE_ID_KEYWORDS):
            mapping["article_id"] = col_name
        elif not mapping.get("document_name") and _match(col_name, _DOC_NAME_KEYWORDS):
            mapping["document_name"] = col_name
        elif not mapping.get("section") and _match(col_name, _SECTION_KEYWORDS):
            mapping["section"] = col_name
        elif not mapping.get("additional_context") and _match(col_name, _CONTEXT_KEYWORDS):
            mapping["additional_context"] = col_name

    if mapping.get("article_text"):
        # No document-name column matched → infer it from repeated short values.
        if not mapping.get("document_name"):
            doc_col = _find_document_column_by_content(
                preview, exclude={_col_name_to_letter(headers, mapping["article_text"])}
            )
            if doc_col:
                mapping["document_name"] = doc_col
                logger.info("制度名列按内容推断: '%s' → document_name", doc_col)
        logger.info(
            "启发式列检测: article_text='%s', mapping=%s",
            mapping["article_text"], {k: v for k, v in mapping.items() if k != "article_text"},
        )
        return mapping

    return _find_text_column_by_content(preview)


def _col_name_to_letter(headers: dict[str, str], col_name: str) -> str:
    """表头名 → 列字母(反向查找)。"""
    for letter, name in headers.items():
        if name == col_name:
            return letter
    return ""


def _find_text_column_by_content(preview: dict) -> dict[str, str] | None:
    """兜底 1:平均文本长度最长的列(>20 字符)当作条款内容列。"""
    samples = preview.get("samples", [])
    headers = preview.get("column_headers", {})
    if not samples or not headers:
        return None

    col_lengths: dict[str, list[int]] = {}
    for s in samples:
        for col_letter, value in s.get("columns", {}).items():
            col_lengths.setdefault(col_letter, []).append(len(value))

    best_col = None
    best_avg = 0
    for col_letter, lengths in col_lengths.items():
        if _match(headers.get(col_letter, ""), _FULL_TEXT_KEYWORDS):
            continue
        avg = sum(lengths) / len(lengths) if lengths else 0
        if avg > best_avg and avg > 20:
            best_avg = avg
            best_col = col_letter

    if best_col and best_col in headers:
        mapping: dict[str, str] = {"article_text": headers[best_col]}
        doc_col = _find_document_column_by_content(preview, exclude={best_col})
        if doc_col:
            mapping["document_name"] = doc_col
            logger.info("制度名列按内容推断: '%s' → document_name", doc_col)
        logger.info(
            "内容分析回退: 最长平均文本列 '%s' (%s) → article_text (avg %d chars)",
            best_col, headers[best_col], int(best_avg),
        )
        return mapping
    return None


def _find_document_column_by_content(
    preview: dict, exclude: set[str] | None = None,
) -> str | None:
    """兜底:按"短值 + 高重复率"推断制度名列。"""
    samples = preview.get("samples", [])
    headers = preview.get("column_headers", {})
    if not samples or not headers:
        return None
    exclude = exclude or set()

    stats: dict[str, dict] = {}
    for s in samples:
        for col_letter, value in s.get("columns", {}).items():
            if not value or col_letter in exclude:
                continue
            st = stats.setdefault(col_letter, {"rows": 0, "vals": set(), "max_len": 0})
            st["rows"] += 1
            st["vals"].add(value)
            st["max_len"] = max(st["max_len"], len(value))

    best_letter, best_score = None, 0.0
    for col_letter, st in stats.items():
        if st["rows"] < 3 or len(st["vals"]) <= 1:
            continue
        repetition = st["rows"] / len(st["vals"])
        length_weight = 1.0 if 2 <= st["max_len"] <= 40 else 0.1
        score = repetition * length_weight
        if score > best_score:
            best_score, best_letter = score, col_letter

    if best_letter and best_letter in headers:
        return headers[best_letter]
    return None


def _normalize_column_header(name: str | None) -> str:
    """归一化表头:去括号注释、特殊字符、空白。"条款内容（原文）" → "条款内容"。"""
    if name is None:
        return ""
    text = str(name).strip()
    text = re.sub(r"[（(][^）)]*[）)]", "", text)
    text = re.sub(r"[【】〈〉《》\s]+", "", text)
    return text.strip()


def _normalize_mapping(mapping: dict[str, str]) -> dict[str, str]:
    return {k: _normalize_column_header(v) for k, v in mapping.items()}


def _infer_article_id_from_text(text: str) -> str:
    """从条款文本推断条款号(第X条/第X节/第X章/Article X/编号项)。"""
    if not text or not text.strip():
        return "空白条款"

    patterns = [
        r"第([一二三四五六七八九十百零\d]+)条",
        r"第([一二三四五六七八九十百零\d]+)节",
        r"第([一二三四五六七八九十百零\d]+)章",
        r"(Article\s+[IVXLCDM\d]+)",
        r"^(\d+[\.\、\)）]\s*)",
    ]

    for pattern in patterns:
        m = re.search(pattern, text)
        if m:
            return m.group(0).strip()

    snippet = text.strip()[:30]
    return f"[{snippet}...]" if len(text) > 30 else snippet


def read_excel_preview(file_path: str | Path, max_rows: int = 30) -> dict[str, Any]:
    """读表头 + 样本行,供列结构分析。

    自动定位真实表头行,并在多个 sheet 中选最像数据表的一个(写入 ``active_sheet``)。
    """
    import openpyxl

    path = Path(file_path)
    result: dict[str, Any] = {
        "file_path": str(path), "sheet_names": [], "active_sheet": "",
        "total_rows": 0, "column_headers": {}, "samples": [], "error": None,
    }

    if not path.exists():
        result["error"] = f"文件不存在: {file_path}"
        return result

    try:
        # read_only mode misreads files with incomplete dimension metadata.
        wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    except Exception as e:
        result["error"] = f"无法打开 Excel 文件: {e}"
        return result

    result["sheet_names"] = wb.sheetnames
    if not wb.sheetnames:
        result["error"] = "Excel 文件中没有工作表"
        wb.close()
        return result

    # Pick the sheet with the most header-like row; fall back to the first.
    best_sheet = wb.sheetnames[0]
    best_score = (-1, -1)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        hr = _find_header_row(ws)
        header_values = next(ws.iter_rows(min_row=hr, max_row=hr, values_only=True), ())
        score = _header_row_hits(header_values)
        if score > best_score:
            best_sheet, best_score = sheet, score

    result["active_sheet"] = best_sheet
    ws = wb[best_sheet]

    header_row_idx = _find_header_row(ws)
    header_row = next(
        ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True), None
    )
    col_letters: dict[int, str] = {}
    if header_row:
        for idx, val in enumerate(header_row):
            col_letters[idx] = _col_idx_to_letter(idx)
            result["column_headers"][col_letters[idx]] = str(val).strip() if val is not None else ""

    sample_count = 0
    data_start = header_row_idx + 1
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=data_start, values_only=True), start=data_start
    ):
        if sample_count >= max_rows:
            break
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        row_data = {
            col_letters.get(col_idx, _col_idx_to_letter(col_idx)): (str(val).strip() if val is not None else "")
            for col_idx, val in enumerate(row)
        }
        result["samples"].append({"row": row_idx, "columns": row_data})
        sample_count += 1

    total = 0
    for row in ws.iter_rows(min_row=data_start, values_only=True):
        if any(v is not None and str(v).strip() != "" for v in row):
            total += 1
        if total >= 20000:
            break
    result["total_rows"] = total

    wb.close()
    logger.info(
        "Excel preview: %s [%s] header_row=%d — %d samples / ~%d total rows, %d headers",
        path.name, best_sheet, header_row_idx, sample_count, total, len(result["column_headers"]),
    )
    return result


def extract_excel_articles(
    file_path: str | Path,
    column_mapping: dict[str, str],
    sheet_name: str | None = None,
    max_rows: int = 20000,
) -> dict[str, Any]:
    """按列映射抽取 Excel 全部行,按 document_name 分组。"""
    import openpyxl

    path = Path(file_path)
    result: dict[str, Any] = {
        "articles": [], "document_groups": {}, "total_rows": 0,
        "unique_documents": [], "error": None,
    }

    if not path.exists():
        result["error"] = f"文件不存在: {file_path}"
        return result

    norm_mapping = _normalize_mapping(column_mapping)

    try:
        # read_only mode misreads files with incomplete dimension metadata.
        wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    except Exception as e:
        result["error"] = f"无法打开 Excel 文件: {e}"
        return result

    target_sheet = sheet_name or wb.sheetnames[0]
    if target_sheet not in wb.sheetnames:
        result["error"] = f"工作表 '{target_sheet}' 不存在。可用工作表: {', '.join(wb.sheetnames)}"
        wb.close()
        return result
    ws = wb[target_sheet]

    header_row_idx = _find_header_row(ws)
    header_row = next(
        ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True), None
    )
    if not header_row:
        result["error"] = f"Excel 第 {header_row_idx} 行为空（需要表头行）"
        wb.close()
        return result

    header_map: dict[str, int] = {}
    for idx, val in enumerate(header_row):
        norm = _normalize_column_header(val)
        if norm:
            header_map[norm] = idx

    col_indices: dict[str, int] = {}
    missing_fields: list[str] = []
    for field_name, norm_target in norm_mapping.items():
        if norm_target in header_map:
            col_indices[field_name] = header_map[norm_target]
        else:
            found = False
            for norm_header, h_idx in header_map.items():
                if norm_target in norm_header or norm_header in norm_target:
                    col_indices[field_name] = h_idx
                    found = True
                    break
            if not found:
                missing_fields.append(f"{field_name} (目标列: {column_mapping.get(field_name, '?')})")

    if "article_text" not in col_indices:
        result["error"] = (
            f"无法在 Excel 中找到以下字段对应的列: {', '.join(missing_fields)}。"
            f"可用的表头: {', '.join(header_map.keys())}"
            if missing_fields else "column_mapping 中缺少必需的 'article_text' 字段"
        )
        wb.close()
        return result

    text_col = col_indices["article_text"]
    id_col = col_indices.get("article_id", -1)
    doc_col = col_indices.get("document_name", -1)
    section_col = col_indices.get("section", -1)
    context_col = col_indices.get("additional_context", -1)

    row_count = 0
    document_groups: dict[str, list[dict[str, Any]]] = {}

    data_start = header_row_idx + 1
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=data_start, values_only=True), start=data_start
    ):
        if row_count >= max_rows:
            logger.warning("Reached max_rows limit (%d) for Excel extraction", max_rows)
            break

        if text_col >= len(row) or row[text_col] is None:
            continue
        article_text = str(row[text_col]).strip()
        if not article_text:
            continue

        article_id = ""
        if id_col >= 0 and id_col < len(row) and row[id_col] is not None:
            article_id = str(row[id_col]).strip()

        document_name = ""
        if doc_col >= 0 and doc_col < len(row) and row[doc_col] is not None:
            document_name = str(row[doc_col]).strip()

        section = ""
        if section_col >= 0 and section_col < len(row) and row[section_col] is not None:
            section = str(row[section_col]).strip()

        full_text_parts: list[str] = []
        if context_col >= 0 and context_col < len(row) and row[context_col] is not None:
            ctx = str(row[context_col]).strip()
            if ctx and ctx != article_text:
                full_text_parts.append(ctx)
        full_text_parts.append(article_text)
        full_text = "\n".join(full_text_parts)

        if not article_id:
            article_id = _infer_article_id_from_text(article_text)

        article_dict = {
            "article_id": article_id,
            "full_text": full_text,
            "document_name": document_name,
            "section": section,
            "row_number": row_idx,
            "source_sheet": target_sheet,
        }
        result["articles"].append(article_dict)
        row_count += 1

        group_key = document_name or "_未命名文档_"
        document_groups.setdefault(group_key, []).append(article_dict)

    wb.close()
    result["total_rows"] = row_count
    result["document_groups"] = document_groups
    result["unique_documents"] = sorted(document_groups.keys())

    logger.info(
        "Excel extraction: %d articles across %d documents from %s",
        row_count, len(document_groups), path.name,
    )
    return result


def _col_idx_to_letter(idx: int) -> str:
    """0 基列号 → Excel 列字母(0→A, 25→Z, 26→AA)。"""
    letters = ""
    n = idx
    while n >= 0:
        letters = chr(ord("A") + (n % 26)) + letters
        n = n // 26 - 1
    return letters
