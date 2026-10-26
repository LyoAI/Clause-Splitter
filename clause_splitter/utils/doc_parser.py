"""本地纯文本抽取(PDF / DOCX / DOC / XLSX / TXT),按内容哈希缓存。"""

from __future__ import annotations

import hashlib
import time
from pathlib import Path


_CACHE_DIR = Path(__file__).resolve().parent.parent.parent / ".cache" / "docs"
_MAX_CACHE_ENTRIES = 100


def extract_text(
    file_path: str | Path,
    max_chars: int = 20000,
    use_cache: bool = True,
) -> str:
    """按后缀抽取纯文本;失败时返回占位字符串。"""
    path = Path(file_path)
    if not path.exists():
        return ""

    if use_cache:
        file_hash = _hash_file(path)
        cached = _read_cache(file_hash)
        if cached is not None:
            return cached[:max_chars]

    suffix = path.suffix.lower()

    try:
        if suffix == ".pdf":
            text = _parse_pdf(path, max_chars)
        elif suffix == ".docx":
            text = _parse_docx(path, max_chars)
        elif suffix == ".doc":
            text = _parse_doc(path, max_chars)
        elif suffix in (".xlsx", ".xls"):
            text = _parse_xlsx(path, max_chars)
        elif suffix in (".png", ".jpg", ".jpeg", ".bmp", ".tiff"):
            text = f"[Image] {path.name} — OCR not available, manual review required."
        else:
            text = path.read_text(encoding="utf-8", errors="replace")[:max_chars]
    except Exception as e:
        text = f"[Parse error: {e}]"

    if use_cache and text and not text.startswith("[Parse error") and not text.startswith("[Image]"):
        _write_cache(file_hash, text, path.name)

    return text


def clear_cache(older_than_days: float = 0) -> int:
    """清理解析缓存;返回移除条目数(0 天 = 全部)。"""
    if not _CACHE_DIR.exists():
        return 0

    removed = 0
    cutoff = time.time() - older_than_days * 86400 if older_than_days > 0 else float("inf")

    for cache_file in _CACHE_DIR.glob("*.txt"):
        try:
            if cache_file.stat().st_mtime < cutoff:
                cache_file.unlink()
                removed += 1
        except OSError:
            pass

    return removed


def _hash_file(path: Path) -> str:
    """流式计算文件 SHA256。"""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(8192):
            h.update(chunk)
    return h.hexdigest()


def _read_cache(file_hash: str) -> str | None:
    cache_path = _CACHE_DIR / f"{file_hash}.txt"
    if not cache_path.exists():
        return None
    try:
        content = cache_path.read_text(encoding="utf-8")
        if content.startswith("# cached:"):
            _, body = content.split("\n", 1)
            return body
        return content
    except Exception:
        return None


def _write_cache(file_hash: str, text: str, original_name: str) -> None:
    _CACHE_DIR.mkdir(parents=True, exist_ok=True)
    header = f"# cached: {original_name} @ {time.strftime('%Y-%m-%d %H:%M:%S')}\n"
    cache_path = _CACHE_DIR / f"{file_hash}.txt"
    cache_path.write_text(header + text, encoding="utf-8")

    entries = sorted(_CACHE_DIR.glob("*.txt"), key=lambda p: p.stat().st_mtime)
    if len(entries) > _MAX_CACHE_ENTRIES:
        for old in entries[:len(entries) - _MAX_CACHE_ENTRIES]:
            try:
                old.unlink()
            except OSError:
                pass


def _parse_pdf(path: Path, max_chars: int) -> str:
    """pdfplumber 抽取 PDF 文本。"""
    import pdfplumber
    texts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                texts.append(t)
            if sum(len(x) for x in texts) >= max_chars:
                break
    return "\n".join(texts)[:max_chars]


def _parse_docx(path: Path, max_chars: int) -> str:
    """python-docx 抽取 .docx 文本。"""
    from docx import Document
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)[:max_chars]


def _parse_doc(path: Path, max_chars: int) -> str:
    """旧版 .doc:先试 antiword,再按 .docx 解析(处理误命名文件)。"""
    import subprocess

    try:
        result = subprocess.run(
            ["antiword", str(path)],
            capture_output=True, text=True, timeout=30,
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout.strip()[:max_chars]
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass

    try:
        from docx import Document
        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        if paragraphs:
            return "\n".join(paragraphs)[:max_chars]
    except Exception:
        pass

    return (
        f"[Legacy .doc] {path.name} — unable to extract text. "
        f"Please convert to .docx or .pdf."
    )


def _parse_xlsx(path: Path, max_chars: int) -> str:
    """openpyxl 抽取 Excel 文本(所有 sheet,tab 分隔)。"""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    lines: list[str] = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        lines.append(f"[Sheet: {sheet}]")
        for row in ws.iter_rows(values_only=True):
            row_texts = [str(v) for v in row if v is not None]
            if row_texts:
                lines.append("\t".join(row_texts))
    return "\n".join(lines)[:max_chars]
